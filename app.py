import io
import json
import os
import re
import time
from typing import Any, Dict, List

import streamlit as st
from google import genai
from google.genai import types
from pypdf import PdfReader
from docx import Document


# ============================================================
# APP CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="ATS Resume Analyzer",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)

MODEL_CANDIDATES = [
    "gemini-3.6-flash",
    "gemini-3.5-flash",
    "gemini-3.5-flash-lite",
]
MAX_RETRIES_PER_MODEL = 3
RETRY_DELAYS_SECONDS = [2, 4, 8]


# ============================================================
# STYLING
# ============================================================

st.markdown(
    """
    <style>
        .main-title {
            font-size: 2.7rem;
            font-weight: 800;
            margin-bottom: 0.2rem;
        }

        .subtitle {
            color: #666;
            font-size: 1.05rem;
            margin-bottom: 1.5rem;
        }

        .score-card {
            padding: 1.2rem;
            border-radius: 14px;
            border: 1px solid rgba(128,128,128,0.25);
            text-align: center;
        }

        .score-number {
            font-size: 3.5rem;
            font-weight: 800;
            line-height: 1;
        }

        .small-label {
            color: #777;
            font-size: 0.9rem;
        }

        .keyword {
            display: inline-block;
            padding: 0.35rem 0.65rem;
            margin: 0.2rem;
            border-radius: 999px;
            background: rgba(255, 165, 0, 0.14);
            border: 1px solid rgba(255, 165, 0, 0.35);
            font-size: 0.88rem;
        }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# API KEY
# ============================================================

def get_gemini_api_key() -> str:
    """Read the Gemini API key from Streamlit secrets or environment."""
    try:
        value = st.secrets.get("GEMINI_API_KEY", "")
        if value:
            return str(value).strip()
    except Exception:
        pass

    return os.getenv("GEMINI_API_KEY", "").strip()


# ============================================================
# FILE EXTRACTION
# ============================================================

def extract_pdf_text(data: bytes) -> str:
    """Extract text from a text-based PDF."""
    reader = PdfReader(io.BytesIO(data))
    pages = []

    for page in reader.pages:
        pages.append(page.extract_text() or "")

    return "\n".join(pages).strip()


def extract_docx_text(data: bytes) -> str:
    """Extract text from DOCX paragraphs and tables."""
    document = Document(io.BytesIO(data))
    parts: List[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def extract_resume_text(uploaded_file) -> str:
    """Extract resume text from PDF, DOCX, or TXT."""
    filename = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        return extract_pdf_text(data)

    if filename.endswith(".docx"):
        return extract_docx_text(data)

    if filename.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()

    raise ValueError(
        "Unsupported file type. Please upload a PDF, DOCX, or TXT file."
    )


# ============================================================
# JSON HELPERS
# ============================================================

def strip_code_fences(text: str) -> str:
    """Remove Markdown code fences around JSON if Gemini adds them."""
    cleaned = text.strip()

    cleaned = re.sub(
        r"^```(?:json)?\s*",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )

    cleaned = re.sub(r"\s*```$", "", cleaned)

    return cleaned.strip()


def parse_json_response(text: str) -> Dict[str, Any]:
    """Parse Gemini's JSON response robustly."""
    cleaned = strip_code_fences(text)

    try:
        parsed = json.loads(cleaned)
        if isinstance(parsed, dict):
            return parsed
    except json.JSONDecodeError:
        pass

    # Fallback: find the outermost JSON object.
    start = cleaned.find("{")
    end = cleaned.rfind("}")

    if start != -1 and end > start:
        try:
            parsed = json.loads(cleaned[start : end + 1])
            if isinstance(parsed, dict):
                return parsed
        except json.JSONDecodeError:
            pass

    raise ValueError("Gemini returned an invalid JSON response.")


# ============================================================
# GEMINI ANALYSIS
# ============================================================

def build_analysis_prompt(resume_text: str, job_description: str) -> str:
    """Build the ATS analysis prompt."""
    if job_description.strip():
        job_section = job_description.strip()
        job_mode = (
            "A job description was supplied. Perform job-specific "
            "keyword and relevance matching."
        )
    else:
        job_section = (
            "[No job description supplied. Perform only a general "
            "ATS-readiness analysis. Do not claim exact job-specific "
            "keyword matching.]"
        )
        job_mode = (
            "No job description was supplied. Focus on general ATS "
            "readiness and clearly state that keyword matching is general."
        )

    return f"""
You are an expert ATS resume auditor and professional career coach.

Analyze the resume using ONLY the information present in the resume.
Do not invent qualifications, experience, achievements, certifications,
skills, employers, dates, or education.

IMPORTANT:
An ATS score is not universal. Different applicant tracking systems
use different ranking methods. Therefore, the score you produce is an
AI-based ATS-readiness estimate from 0 to 100, not an official score
from a specific ATS vendor.

{job_mode}

================ JOB DESCRIPTION ================
{job_section}

================ RESUME TEXT ================
{resume_text}

================ ANALYSIS REQUIREMENTS ================

Evaluate:

1. Keyword alignment
2. Resume structure and standard section headings
3. Experience relevance
4. Skills section
5. Clarity and readability
6. Contact information and essential sections
7. Action verbs and measurable achievements
8. ATS risks that can be inferred from the extracted text

Common ATS risks may include:
- unusual section headings
- missing contact information
- weak or generic keywords
- vague bullet points
- lack of measurable achievements
- poor chronology
- excessive or unclear text
- suspiciously repetitive keywords
- information that may be difficult for an ATS to interpret

Do NOT claim that the resume contains columns, graphics, icons,
tables, headers, footers, text boxes, or visual formatting problems
unless that issue can actually be established from the extracted
content. If visual inspection is required, say so.

For a supplied job description:
- identify important keywords/skills present in the job description
- identify relevant keywords already present in the resume
- identify important missing keywords
- distinguish between exact/near-exact keyword matches and reasonable
  semantic matches
- do not recommend a keyword merely because it sounds useful;
  prioritize terms actually supported by the job description

Scoring:
- ats_score: integer 0-100
- category scores: integers 0-100
- Keep the score realistic. Do not automatically give a high score.
- Do not penalize the user for information that cannot be verified
  from extracted text.

Recommendations must be actionable and concise.
Do not rewrite the entire resume.

================ REQUIRED JSON ================

Return ONLY valid JSON with this exact structure:

{{
  "ats_score": 0,
  "score_summary": "Short explanation of the overall score.",
  "category_scores": {{
    "keyword_match": 0,
    "format_and_structure": 0,
    "experience_relevance": 0,
    "skills": 0,
    "clarity_and_readability": 0,
    "contact_and_sections": 0
  }},
  "strengths": [
    "Strength 1",
    "Strength 2",
    "Strength 3"
  ],
  "critical_issues": [
    "Critical issue 1",
    "Critical issue 2"
  ],
  "improvements": [
    {{
      "priority": "High",
      "issue": "Specific issue",
      "recommendation": "Specific action the user should take",
      "example": "Short example of how to improve it"
    }}
  ],
  "missing_keywords": [
    "keyword 1",
    "keyword 2"
  ],
  "matched_keywords": [
    "keyword 1",
    "keyword 2"
  ],
  "section_feedback": [
    {{
      "section": "Experience",
      "status": "Good",
      "feedback": "Specific feedback"
    }}
  ],
  "ats_checklist": {{
    "standard_headings": true,
    "keyword_alignment": true,
    "readable_contact_information": true,
    "action_verbs_and_metrics": true,
    "clear_chronology": true
  }}
}}
"""


def is_retryable_error(exc: Exception) -> bool:
    """Return True for temporary Gemini capacity/rate-limit errors."""
    message = str(exc).upper()
    retryable_markers = [
        "503",
        "UNAVAILABLE",
        "429",
        "RESOURCE_EXHAUSTED",
        "TOO MANY REQUESTS",
        "DEADLINE EXCEEDED",
        "504",
        "TIMEOUT",
    ]
    return any(marker in message for marker in retryable_markers)


def analyze_resume(
    resume_text: str,
    job_description: str,
) -> tuple[Dict[str, Any], str]:
    """
    Analyze the resume with Gemini using retries and automatic model fallback.

    Temporary 503/429/time-out errors are retried with exponential backoff.
    If a model remains unavailable, the next Flash model is tried.
    """
    api_key = get_gemini_api_key()

    if not api_key:
        raise RuntimeError(
            "Gemini API key is missing. Add GEMINI_API_KEY to "
            ".streamlit/secrets.toml or your environment variables."
        )

    client = genai.Client(api_key=api_key)
    prompt = build_analysis_prompt(
        resume_text=resume_text,
        job_description=job_description,
    )

    last_error = None

    for model_name in MODEL_CANDIDATES:
        for attempt in range(MAX_RETRIES_PER_MODEL):
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        response_mime_type="application/json",
                    ),
                )

                if not response.text:
                    raise RuntimeError(
                        f"{model_name} returned an empty response."
                    )

                return parse_json_response(response.text), model_name

            except Exception as exc:
                last_error = exc

                if not is_retryable_error(exc):
                    raise RuntimeError(
                        f"Gemini request failed with {model_name}: {exc}"
                    ) from exc

                if attempt < MAX_RETRIES_PER_MODEL - 1:
                    time.sleep(RETRY_DELAYS_SECONDS[attempt])

    raise RuntimeError(
        "Gemini is temporarily busy or unavailable. The app tried "
        "multiple Gemini Flash models and automatic retries. Please wait "
        "a few minutes and try again."
    ) from last_error


# ============================================================
# VALIDATION / NORMALIZATION
# ============================================================

def safe_score(value: Any, default: int = 0) -> int:
    """Convert an arbitrary score to an integer between 0 and 100."""
    try:
        score = int(float(value))
        return max(0, min(100, score))
    except (TypeError, ValueError):
        return default


def normalize_analysis(data: Dict[str, Any]) -> Dict[str, Any]:
    """Make the model output safe for rendering."""
    category_defaults = {
        "keyword_match": 0,
        "format_and_structure": 0,
        "experience_relevance": 0,
        "skills": 0,
        "clarity_and_readability": 0,
        "contact_and_sections": 0,
    }

    categories = data.get("category_scores", {})
    if not isinstance(categories, dict):
        categories = {}

    normalized_categories = {
        key: safe_score(categories.get(key, default))
        for key, default in category_defaults.items()
    }

    normalized = {
        "ats_score": safe_score(data.get("ats_score", 0)),
        "score_summary": str(data.get("score_summary", "")).strip(),
        "category_scores": normalized_categories,
        "strengths": data.get("strengths", []),
        "critical_issues": data.get("critical_issues", []),
        "improvements": data.get("improvements", []),
        "missing_keywords": data.get("missing_keywords", []),
        "matched_keywords": data.get("matched_keywords", []),
        "section_feedback": data.get("section_feedback", []),
        "ats_checklist": data.get("ats_checklist", {}),
    }

    # Ensure list-like fields are actually lists.
    for key in [
        "strengths",
        "critical_issues",
        "missing_keywords",
        "matched_keywords",
        "improvements",
        "section_feedback",
    ]:
        if not isinstance(normalized[key], list):
            normalized[key] = []

    if not isinstance(normalized["ats_checklist"], dict):
        normalized["ats_checklist"] = {}

    return normalized


# ============================================================
# DISPLAY HELPERS
# ============================================================

def score_label(score: int) -> str:
    if score >= 80:
        return "🟢 Strong ATS readiness"
    if score >= 60:
        return "🟡 Moderate ATS readiness"
    return "🔴 Needs improvement"


def render_keyword_list(
    keywords: List[Any],
    empty_message: str,
) -> None:
    """Render keywords as compact pills."""
    clean_keywords = [
        str(keyword).strip()
        for keyword in keywords
        if str(keyword).strip()
    ]

    if not clean_keywords:
        st.info(empty_message)
        return

    html = " ".join(
        f'<span class="keyword">{keyword}</span>'
        for keyword in clean_keywords
    )

    st.markdown(html, unsafe_allow_html=True)


def render_bullet_list(items: List[Any], empty_message: str) -> None:
    """Render a list as Streamlit bullets."""
    clean_items = [
        str(item).strip()
        for item in items
        if str(item).strip()
    ]

    if not clean_items:
        st.info(empty_message)
        return

    for item in clean_items:
        st.write(f"• {item}")


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:
    st.header("📌 How it works")

    st.write("**1. Upload** your resume.")
    st.write("**2. Add** a target job description if available.")
    st.write("**3. Analyze** with Gemini Flash.")
    st.write("**4. Review** your ATS-readiness score.")
    st.write("**5. Improve** the resume using the recommendations.")

    st.divider()

    st.subheader("Supported files")
    st.write("• PDF")
    st.write("• DOCX")
    st.write("• TXT")

    st.divider()

    st.warning(
        "Privacy: your resume is sent to the Gemini API for analysis. "
        "Do not upload sensitive documents unless you are comfortable "
        "with that processing."
    )


# ============================================================
# MAIN HEADER
# ============================================================

st.markdown(
    '<div class="main-title">📄 ATS Resume Analyzer</div>',
    unsafe_allow_html=True,
)

st.markdown(
    '<div class="subtitle">'
    "Get an AI-based ATS-readiness score, keyword analysis, "
    "resume strengths, issues, and actionable improvements."
    "</div>",
    unsafe_allow_html=True,
)


# ============================================================
# INPUT AREA
# ============================================================

upload_col, jd_col = st.columns([1, 1])

with upload_col:
    uploaded_file = st.file_uploader(
        "Upload your resume",
        type=["pdf", "docx", "txt"],
        help="Upload a text-based PDF, DOCX, or TXT resume.",
    )

with jd_col:
    job_description = st.text_area(
        "Target job description (optional)",
        height=150,
        placeholder=(
            "Paste the job description here. "
            "Adding it makes keyword and relevance analysis more useful."
        ),
    )


# ============================================================
# ANALYZE
# ============================================================

if uploaded_file is not None:
    st.success(f"Ready to analyze: **{uploaded_file.name}**")

    if st.button(
        "🔍 Analyze Resume",
        type="primary",
        use_container_width=True,
    ):
        try:
            with st.spinner(
                "Extracting resume text and analyzing it with Gemini..."
            ):
                resume_text = extract_resume_text(uploaded_file)

                if not resume_text:
                    st.error(
                        "No readable text was found in the uploaded file."
                    )
                    st.info(
                        "If this is a scanned/image-only PDF, use a "
                        "text-based PDF or DOCX version of the resume."
                    )
                    st.stop()

                if len(resume_text) < 80:
                    st.warning(
                        "Very little text was extracted. The analysis "
                        "may not be reliable."
                    )

                # Keep the prompt at a practical size.
                if len(resume_text) > 60000:
                    resume_text = resume_text[:60000]
                    st.warning(
                        "The extracted resume was very long, so the "
                        "analysis input was limited to 60,000 characters."
                    )

                analysis, model_used = analyze_resume(
                    resume_text=resume_text,
                    job_description=job_description,
                )

                st.session_state["analysis"] = normalize_analysis(analysis)
                st.session_state["model_used"] = model_used
                st.session_state["resume_name"] = uploaded_file.name

        except Exception as exc:
            st.error(f"Analysis failed: {exc}")
            st.caption(
                "Check your Gemini API key, internet connection, "
                "model availability, and uploaded file."
            )


# ============================================================
# RESULTS
# ============================================================

if "analysis" in st.session_state:
    result = st.session_state["analysis"]

    st.divider()

    # --------------------------------------------------------
    # OVERALL SCORE
    # --------------------------------------------------------

    st.subheader("🎯 ATS Readiness Score")
    st.caption(f"Gemini model used: `{st.session_state.get('model_used', MODEL_CANDIDATES[0])}`")

    score = result["ats_score"]

    score_col, explanation_col = st.columns([1, 3])

    with score_col:
        st.markdown(
            f"""
            <div class="score-card">
                <div class="small-label">Overall score</div>
                <div class="score-number">{score}</div>
                <div class="small-label">out of 100</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with explanation_col:
        st.progress(score / 100)
        st.write(f"**{score_label(score)}**")
        st.write(result["score_summary"])

    # --------------------------------------------------------
    # CATEGORY SCORES
    # --------------------------------------------------------

    st.subheader("📊 Category Scores")

    category_names = {
        "keyword_match": "Keyword Match",
        "format_and_structure": "Format & Structure",
        "experience_relevance": "Experience Relevance",
        "skills": "Skills",
        "clarity_and_readability": "Clarity & Readability",
        "contact_and_sections": "Contact & Sections",
    }

    category_items = list(result["category_scores"].items())

    for start in range(0, len(category_items), 3):
        cols = st.columns(3)

        for col, (key, value) in zip(
            cols,
            category_items[start : start + 3],
        ):
            with col:
                st.metric(
                    category_names.get(
                        key,
                        key.replace("_", " ").title(),
                    ),
                    f"{value}/100",
                )
                st.progress(value / 100)

    # --------------------------------------------------------
    # STRENGTHS / CRITICAL ISSUES
    # --------------------------------------------------------

    left_col, right_col = st.columns(2)

    with left_col:
        st.subheader("✅ Strengths")
        render_bullet_list(
            result["strengths"],
            "No strengths were returned.",
        )

    with right_col:
        st.subheader("🚨 Critical Issues")
        render_bullet_list(
            result["critical_issues"],
            "No critical issues were identified.",
        )

    # --------------------------------------------------------
    # KEYWORDS
    # --------------------------------------------------------

    st.subheader("🔑 Keyword Analysis")

    keyword_left, keyword_right = st.columns(2)

    with keyword_left:
        st.markdown("**Matched keywords**")
        render_keyword_list(
            result["matched_keywords"],
            "No matched keywords were returned.",
        )

    with keyword_right:
        st.markdown("**Missing / important keywords**")
        render_keyword_list(
            result["missing_keywords"],
            "No major missing keywords were identified.",
        )

    # --------------------------------------------------------
    # IMPROVEMENTS
    # --------------------------------------------------------

    st.subheader("🛠️ Recommended Improvements")

    improvements = result["improvements"]

    if not improvements:
        st.info("No specific improvement recommendations were returned.")
    else:
        for index, item in enumerate(improvements, start=1):
            if not isinstance(item, dict):
                st.write(f"• {item}")
                continue

            priority = str(item.get("priority", "Medium"))
            issue = str(item.get("issue", "Improvement"))

            with st.expander(
                f"{index}. [{priority}] {issue}",
                expanded=(index == 1),
            ):
                recommendation = str(
                    item.get("recommendation", "")
                ).strip()

                example = str(
                    item.get("example", "")
                ).strip()

                if recommendation:
                    st.markdown(
                        f"**Recommendation:** {recommendation}"
                    )

                if example:
                    st.markdown(
                        f"**Example:** {example}"
                    )

    # --------------------------------------------------------
    # SECTION FEEDBACK
    # --------------------------------------------------------

    st.subheader("📌 Section-by-Section Feedback")

    feedback = result["section_feedback"]

    if not feedback:
        st.info("No section-specific feedback was returned.")
    else:
        for item in feedback:
            if not isinstance(item, dict):
                st.write(f"• {item}")
                continue

            section = str(item.get("section", "Section"))
            status = str(item.get("status", "Review"))
            text = str(item.get("feedback", "")).strip()

            st.markdown(f"**{section} — {status}**")
            if text:
                st.write(text)

    # --------------------------------------------------------
    # ATS CHECKLIST
    # --------------------------------------------------------

    st.subheader("☑️ ATS Checklist")

    checklist = result["ats_checklist"]

    checklist_labels = {
        "standard_headings": "Standard section headings",
        "keyword_alignment": "Keyword alignment",
        "readable_contact_information": "Readable contact information",
        "action_verbs_and_metrics": "Action verbs and measurable results",
        "clear_chronology": "Clear chronology",
    }

    if not checklist:
        st.info("No checklist data was returned.")
    else:
        checklist_cols = st.columns(2)

        for index, (key, value) in enumerate(checklist.items()):
            label = checklist_labels.get(
                key,
                key.replace("_", " ").title(),
            )

            icon = "✅" if bool(value) else "⚠️"

            with checklist_cols[index % 2]:
                st.write(f"{icon} **{label}**")

    # --------------------------------------------------------
    # DISCLAIMER
    # --------------------------------------------------------

    st.divider()

    st.caption(
        "Important: ATS scoring varies between applicant-tracking "
        "systems. This application provides an AI-based ATS-readiness "
        "estimate and practical recommendations; it is not an official "
        "score from a specific ATS provider."
    )
